#!/usr/bin/env python3
"""Audit and repair historical meal counts for one customer import batch.

Dry-run is the default. Apply requires a confirmed workbook SHA-256, an exact
expected row count, and matching old order values. The script never deletes data.
"""

import argparse
import calendar
import hashlib
import json
import os
import re
import shlex
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path

import pymysql
from openpyxl import load_workbook


ORDER_PREFIX = re.compile(r"^ORD\d{8}$")
MEAL_TYPES = ("BREAKFAST", "LUNCH", "DINNER")
PACKAGE_NAMES = {"A": "月子餐", "B": "孕期餐", "C": "小月子餐", "D": "营养餐", "F": "营养餐"}


def exact_int(value, label):
    """Return an integral cell value or fail with its field label."""
    if isinstance(value, bool) or not isinstance(value, (int, float)) or int(value) != value:
        raise ValueError("{} must be an integer".format(label))
    return int(value)


def source_counts(path, import_date):
    """Read the month sheet and group lunch/dinner counts by effective customer code."""
    sheet_name = "{}年{}月".format(import_date.year % 100, import_date.month)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        day_count = calendar.monthrange(import_date.year, import_date.month)[1]
        grouped = defaultdict(lambda: {
            "total": 0, "sheet_remaining": 0, "future": 0, "dates": set(),
            "descriptions": [], "phones": set(), "soups": [],
        })
        effective_code = None
        for row_number, row in enumerate(
                sheet.iter_rows(min_row=4, max_col=10 + day_count * 3, values_only=True), start=4):
            code_a, code_b = row[0], row[1]
            if code_b not in (None, ""):
                effective_code = str(code_b).strip()
            elif code_a not in (None, ""):
                effective_code = str(code_a).strip()
            description = row[7]
            if not effective_code or not isinstance(description, str):
                continue
            if "早餐" in description and "午餐" not in description and "晚餐" not in description:
                continue
            if not any(word in description for word in ("午餐", "晚餐", "等通知")):
                continue
            record = grouped[effective_code]
            record["descriptions"].append(description)
            if row[2] not in (None, ""):
                record["phones"].add(re.sub(r"\D", "", str(row[2])))
            if row[6] not in (None, ""):
                record["soups"].append(str(row[6]).strip())
            try:
                record["total"] += exact_int(row[8], "row {} meal count".format(row_number))
                record["sheet_remaining"] += exact_int(row[9], "row {} remaining".format(row_number))
                for day in range(1, day_count + 1):
                    delivery_date = date(import_date.year, import_date.month, day)
                    if delivery_date <= import_date:
                        continue
                    for meal_offset in (1, 2):
                        value = row[10 + (day - 1) * 3 + meal_offset]
                        if value in (None, ""):
                            continue
                        quantity = exact_int(value, "row {} day {}".format(row_number, day))
                        if quantity > 0:
                            record["future"] += quantity
                            record["dates"].add((delivery_date.isoformat(), MEAL_TYPES[meal_offset]))
            except ValueError as exc:
                record["error"] = str(exc)
        for code, record in grouped.items():
            record["available"] = record["sheet_remaining"] + record["future"]
            record["historical_verified"] = record["total"] - record["available"]
            if record["historical_verified"] < 0:
                record["error"] = "source meal count is below available count"
        return grouped
    finally:
        workbook.close()


def db_config(env_path):
    """Read only DB_* settings from the private shell-style environment file."""
    values = {}
    if env_path:
        for line in Path(env_path).read_text().splitlines():
            line = line.strip()
            if not line.startswith("export DB_") or "=" not in line:
                continue
            key, raw = line[len("export "):].split("=", 1)
            tokens = shlex.split(raw, comments=False)
            if tokens:
                values[key] = tokens[0]
    values.update({key: value for key, value in os.environ.items() if key.startswith("DB_")})
    required = ("DB_HOST", "DB_NAME", "DB_USER", "DB_PWD")
    missing = [key for key in required if not values.get(key)]
    if missing:
        raise ValueError("missing database settings: {}".format(", ".join(missing)))
    return values


def delivery_cells(raw):
    """Return the saved future date and meal-type pairs without personal data."""
    if not raw:
        return set()
    values = json.loads(raw) if isinstance(raw, str) else raw
    result = set()
    for item in values:
        if isinstance(item, str):
            raise ValueError("old date-only delivery data cannot be matched safely")
        for meal_type in item.get("mealTypes") or []:
            result.add((item.get("date"), meal_type))
    return result


def inspect_orders(cursor, prefix, customer_code, has_imported_column):
    """Read zero-amount orders from the specified import batch."""
    imported_column = "o.imported_verified_count" if has_imported_column else "0 AS imported_verified_count"
    sql = """
        SELECT o.id, o.order_code, cp.customer_code, o.start_date,
               o.breakfast_count, o.lunch_dinner_count, o.verified_count,
               {}, o.remaining_count, o.status,
               o.delivery_dates
        FROM customer_order o
        JOIN customer_profile cp ON cp.id = o.customer_id
        WHERE o.order_code LIKE %s
          AND o.breakfast_count = 0
          AND o.total_amount = 0
          AND o.final_amount = 0
    """.format(imported_column)
    params = [prefix + "%"]
    if customer_code:
        sql += " AND cp.customer_code = %s"
        params.append(customer_code)
    sql += " ORDER BY o.id"
    cursor.execute(sql, params)
    return cursor.fetchall()


def inspect_missing_profiles(cursor, import_date, customer_code):
    """Read import-day profiles that have no order, for historical-only repair."""
    sql = """
        SELECT cp.id, cp.customer_code, cp.phone, cp.remark,
               cp.special_requirements, cp.create_by, cp.create_time
        FROM customer_profile cp
        LEFT JOIN customer_order o ON o.customer_id = cp.id
        WHERE cp.create_time >= %s AND cp.create_time < %s AND o.id IS NULL
    """
    params = [import_date, import_date + timedelta(days=1)]
    if customer_code:
        sql += " AND cp.customer_code = %s"
        params.append(customer_code)
    sql += " ORDER BY cp.id"
    cursor.execute(sql, params)
    return cursor.fetchall()


def load_parent_packages(cursor):
    """Load active numbered package pools without customer data."""
    cursor.execute("""
        SELECT id, package_name, pool_prefix, pool_start, pool_end
        FROM parent_package WHERE status = 1
    """)
    return cursor.fetchall()


def matching_parent_package(code, packages):
    """Match the importer's unique parent-package name and number-pool rule."""
    match = re.fullmatch(r"([A-Za-z]+)(\d+)", code)
    if not match:
        return None
    expected_name = PACKAGE_NAMES.get(match.group(1).upper())
    hits = []
    for package in packages:
        prefix = package["pool_prefix"]
        if (not prefix or package["pool_start"] is None or package["pool_end"] is None
                or package["package_name"].strip() != expected_name or not code.startswith(prefix)):
            continue
        suffix = code[len(prefix):]
        if suffix.isdigit() and package["pool_start"] <= int(suffix) <= package["pool_end"]:
            hits.append(package["id"])
    return hits[0] if len(hits) == 1 else None


def missing_order_candidate(profile, source, packages, import_date):
    """Validate a profile-only import before proposing its historical order."""
    code = profile["customer_code"]
    values = source.get(code)
    if not values:
        return None, "not present in workbook"
    if values.get("error"):
        return None, values["error"]
    if values["total"] <= 0:
        return None, "no source purchase count"
    if values["available"] != 0:
        return None, "source has future or unused meals but profile has no order"
    if len(values["phones"]) != 1 or next(iter(values["phones"])) != profile["phone"]:
        return None, "source phone differs from stored customer phone"
    parent_id = matching_parent_package(code, packages)
    if parent_id is None:
        return None, "parent package number pool is not unique"
    soups = list(dict.fromkeys(values["soups"]))
    if not soups or any(value not in ("含汤", "不含汤") for value in soups):
        return None, "source soup description is invalid"
    descriptions = values["descriptions"]
    wait_notice = any("等通知" in text for text in descriptions)
    lunch = any("午餐" in text for text in descriptions)
    dinner = any("晚餐" in text for text in descriptions)
    meal_type = None if wait_notice else "LUNCH_DINNER" if lunch and dinner else "LUNCH" if lunch else "DINNER" if dinner else None
    if not wait_notice and meal_type is None:
        return None, "source meal type is absent"
    modes = ["SCHEDULE" if "等通知" in text else "DAILY" if "每日" in text
             else "WEEKDAY" if "工作日" in text else "WEEKEND" if "周末" in text else None
             for text in descriptions]
    if None in modes or (not wait_notice and len(set(modes)) > 1):
        return None, "source schedule mode is ambiguous"
    schedule_mode = "SCHEDULE" if wait_notice else modes[0]
    special = profile["special_requirements"] or ""
    paused = wait_notice or "等通知" in (profile["remark"] or "") or "等通知" in special
    return {
        "customer_id": profile["id"], "code": code, "parent_id": parent_id,
        "total": values["total"], "status": 4 if paused else 2,
        "meal_type": meal_type, "schedule_mode": schedule_mode,
        "start_meal_type": "LUNCH" if meal_type == "LUNCH_DINNER" else meal_type,
        "soup_count": 1 if soups[0] == "含汤" else 0,
        "mixed_soup": len(soups) > 1,
        "side_dish_count": 0 if "无副菜" in special else 1,
        "remark": profile["remark"], "create_by": profile["create_by"],
        "deal_time": profile["create_time"], "import_date": import_date,
    }, None


def log_count(cursor, order_id):
    """Count actual, unrefunded verification records for an order."""
    cursor.execute("""
        SELECT COALESCE(SUM(verification_count), 0) AS verification_count
        FROM meal_verification_log
        WHERE order_id = %s AND COALESCE(deleted, 0) = 0 AND COALESCE(is_refunded, 0) = 0
    """, (order_id,))
    return int(cursor.fetchone()["verification_count"])


def past_plan_count(cursor, order_id, import_date):
    """Check that shifting the order start date will not hide generated plans."""
    cursor.execute("""
        SELECT COUNT(*) AS plan_count
        FROM meal_plan_customer mpc
        JOIN meal_plan mp ON mp.id = mpc.meal_plan_id
        WHERE mpc.order_id = %s AND mp.record_date <= %s
          AND COALESCE(mpc.deleted, 0) = 0 AND COALESCE(mp.deleted, 0) = 0
    """, (order_id, import_date))
    return int(cursor.fetchone()["plan_count"])


def candidate(row, source, import_date, cursor):
    """Validate an old imported order before proposing an exact update."""
    code = row["customer_code"]
    if code not in source:
        return None, "not present in workbook"
    values = source[code]
    if values.get("error"):
        return None, values["error"]
    if row["start_date"] != import_date:
        return None, "start date differs"
    if row["status"] not in (1, 4):
        return None, "status changed"
    if row["imported_verified_count"] != 0:
        return None, "already migrated"
    if row["lunch_dinner_count"] != values["available"]:
        return None, "old purchase count differs"
    if delivery_cells(row["delivery_dates"]) != values["dates"]:
        return None, "future delivery dates differ"
    actual_verified = log_count(cursor, row["id"])
    if actual_verified != 0 or row["verified_count"] != 0:
        return None, "post-import verification exists"
    if row["remaining_count"] != values["available"]:
        return None, "old remaining count differs"
    if past_plan_count(cursor, row["id"], import_date):
        return None, "generated plan exists on or before import date"
    if values["historical_verified"] == 0:
        return None, "no count correction needed"
    return {
        "id": row["id"], "order_code": row["order_code"], "code": code,
        "old_count": row["lunch_dinner_count"], "new_count": values["total"],
        "historical_verified": values["historical_verified"],
        "remaining": values["available"],
        "new_start_date": import_date + timedelta(days=1), "status": row["status"],
    }, None


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", required=True)
    parser.add_argument("--import-date", required=True, type=date.fromisoformat)
    parser.add_argument("--order-code-prefix", required=True)
    scope = parser.add_mutually_exclusive_group(required=True)
    scope.add_argument("--customer-code")
    scope.add_argument("--all", action="store_true", help="audit every matching imported order")
    parser.add_argument("--exclude-customer-code", action="append", default=[],
                        help="explicitly leave one source-conflict customer unchanged; repeatable")
    parser.add_argument("--include-missing-orders", action="store_true",
                        help="also audit profiles whose fully historical order was never created")
    parser.add_argument("--env-file", help="private shell-style file containing DB_* variables")
    parser.add_argument("--apply", action="store_true", help="commit the verified changes")
    parser.add_argument("--expected-count", type=int, help="required with --apply")
    parser.add_argument("--expected-insert-count", type=int,
                        help="required with --apply --include-missing-orders")
    parser.add_argument("--expected-sha256", help="required with --apply")
    parser.add_argument("--backup-file", help="required with --apply; created privately before any write")
    args = parser.parse_args()
    if not ORDER_PREFIX.fullmatch(args.order_code_prefix):
        parser.error("--order-code-prefix must be ORDyyyyMMdd")
    if args.order_code_prefix != "ORD" + args.import_date.strftime("%Y%m%d"):
        parser.error("order code prefix and import date differ")
    if args.apply and (args.expected_count is None or not args.expected_sha256 or not args.backup_file):
        parser.error("--apply requires --expected-count, --expected-sha256 and --backup-file")
    if args.apply and args.include_missing_orders and args.expected_insert_count is None:
        parser.error("--apply --include-missing-orders requires --expected-insert-count")
    workbook_hash = hashlib.sha256(Path(args.workbook).read_bytes()).hexdigest()
    if args.expected_sha256 and workbook_hash.lower() != args.expected_sha256.lower():
        parser.error("workbook SHA-256 does not match")
    source = source_counts(args.workbook, args.import_date)
    if args.customer_code and args.customer_code not in source:
        parser.error("customer code is absent from the selected month sheet")
    print("workbook_sha256={} source_customers={}".format(workbook_hash, len(source)))
    settings = db_config(args.env_file)
    connection = pymysql.connect(
        host=settings["DB_HOST"], port=int(settings.get("DB_PORT", 3306)),
        user=settings["DB_USER"], password=settings["DB_PWD"], database=settings["DB_NAME"],
        charset="utf8mb4", cursorclass=pymysql.cursors.DictCursor,
        autocommit=False, connect_timeout=5, read_timeout=10,
    )
    with connection:
        with connection.cursor() as cursor:
            cursor.execute("START TRANSACTION" if args.apply else "START TRANSACTION READ ONLY")
            cursor.execute("""
                SELECT COUNT(*) AS column_count FROM information_schema.columns
                WHERE table_schema = DATABASE() AND table_name = 'customer_order'
                  AND column_name = 'imported_verified_count'
            """)
            has_imported_column = cursor.fetchone()["column_count"] > 0
            if args.apply and not has_imported_column:
                raise ValueError("run customer_order imported_verified_count migration before --apply")
            orders = inspect_orders(cursor, args.order_code_prefix, args.customer_code, has_imported_column)
            changes = []
            skipped = []
            for row in orders:
                if row["customer_code"] in args.exclude_customer_code:
                    change, reason = None, "explicitly excluded"
                else:
                    change, reason = candidate(row, source, args.import_date, cursor)
                if reason:
                    skipped.append((row["order_code"], reason))
                else:
                    changes.append(change)
            print("matching_orders={} candidate_updates={} skipped={}".format(
                len(orders), len(changes), len(skipped)))
            for change in changes:
                print("order={} customer={} lunch_dinner:{}->{} imported_verified:{} remaining:{} start_date:{}".format(
                    change["order_code"], change["code"], change["old_count"], change["new_count"],
                    change["historical_verified"], change["remaining"], change["new_start_date"]))
            for order_code, reason in skipped[:20]:
                print("skipped order={} reason={}".format(order_code, reason))
            missing_changes = []
            missing_skipped = []
            if args.include_missing_orders:
                profiles = inspect_missing_profiles(cursor, args.import_date, args.customer_code)
                packages = load_parent_packages(cursor)
                for profile in profiles:
                    if profile["customer_code"] in args.exclude_customer_code:
                        change, reason = None, "explicitly excluded"
                    else:
                        change, reason = missing_order_candidate(profile, source, packages, args.import_date)
                    if reason:
                        missing_skipped.append((profile["customer_code"], reason))
                    else:
                        missing_changes.append(change)
                print("profiles_without_order={} candidate_inserts={} skipped={}".format(
                    len(profiles), len(missing_changes), len(missing_skipped)))
                for change in missing_changes:
                    print("missing_order customer={} purchase={} imported_verified={} status={} mixed_soup={}".format(
                        change["code"], change["total"], change["total"],
                        change["status"], change["mixed_soup"]))
                for code, reason in missing_skipped[:20]:
                    print("skipped profile={} reason={}".format(code, reason))
            if args.apply:
                blocking_skipped = [(code, reason) for code, reason in skipped
                                    if reason not in ("already migrated", "no count correction needed",
                                                      "explicitly excluded")]
                blocking_missing = [(code, reason) for code, reason in missing_skipped
                                    if reason not in ("not present in workbook", "no source purchase count",
                                                      "explicitly excluded")]
                if (blocking_skipped or blocking_missing or len(changes) != args.expected_count
                        or args.include_missing_orders
                        and len(missing_changes) != args.expected_insert_count):
                    raise ValueError("candidate count or unsafe skipped orders changed; transaction rolled back")
                if missing_changes:
                    cursor.execute("""
                        SELECT order_code FROM customer_order
                        WHERE order_code LIKE %s ORDER BY order_code DESC LIMIT 1 FOR UPDATE
                    """, (args.order_code_prefix + "%",))
                    last = cursor.fetchone()
                    suffix = last["order_code"][len(args.order_code_prefix):] if last else "0"
                    if not suffix.isdigit():
                        raise ValueError("latest order code suffix is not numeric")
                    next_number = int(suffix) + 1
                    for change in missing_changes:
                        change["order_code"] = args.order_code_prefix + "{:03d}".format(next_number)
                        next_number += 1
                backup = {
                    "workbook_sha256": workbook_hash, "import_date": args.import_date.isoformat(),
                    "updates": changes, "inserts": missing_changes,
                }
                backup_fd = os.open(args.backup_file, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
                with os.fdopen(backup_fd, "w", encoding="utf-8") as backup_out:
                    json.dump(backup, backup_out, ensure_ascii=False, indent=2, default=str)
                for change in changes:
                    status = 2 if change["remaining"] == 0 and change["status"] == 1 else change["status"]
                    cursor.execute("""
                        UPDATE customer_order
                        SET lunch_dinner_count = %s, imported_verified_count = %s,
                            verified_count = %s, start_date = %s, status = %s,
                            update_time = NOW()
                        WHERE id = %s AND lunch_dinner_count = %s
                          AND verified_count = 0 AND imported_verified_count = 0
                          AND remaining_count = %s AND start_date = %s AND status = %s
                    """, (change["new_count"], change["historical_verified"],
                          change["historical_verified"], change["new_start_date"], status,
                          change["id"], change["old_count"], change["remaining"],
                          args.import_date, change["status"]))
                    if cursor.rowcount != 1:
                        raise ValueError("order changed during repair; transaction rolled back")
                if missing_changes:
                    for change in missing_changes:
                        cursor.execute("""
                            INSERT INTO customer_order (
                                customer_id, customer_code, parent_package_id, order_code,
                                lunch_dinner_count, total_count, verified_count,
                                imported_verified_count, remaining_count, deal_time,
                                start_date, start_meal_type, pause_effective_date,
                                status, meal_type, schedule_mode, remark, create_by, create_time,
                                main_dish_count, side_dish_count, veg_count, rice_count,
                                soup_count, rice_type
                            )
                            SELECT %s, %s, %s, %s, %s, %s, %s, %s, 0, %s,
                                   %s, %s, %s, %s, %s, %s, %s, %s, %s,
                                   1, %s, 1, 1, %s, '白米饭'
                            FROM customer_profile cp
                            WHERE cp.id = %s AND cp.customer_code = %s
                              AND NOT EXISTS (
                                  SELECT 1 FROM customer_order existing
                                  WHERE existing.customer_id = cp.id
                              )
                        """, (
                            change["customer_id"], change["code"], change["parent_id"], change["order_code"],
                            change["total"], change["total"], change["total"], change["total"],
                            change["deal_time"], change["import_date"] + timedelta(days=1),
                            change["start_meal_type"],
                            change["import_date"] if change["status"] == 4 else None,
                            change["status"], change["meal_type"], change["schedule_mode"],
                            change["remark"], change["create_by"], change["deal_time"],
                            change["side_dish_count"], change["soup_count"],
                            change["customer_id"], change["code"],
                        ))
                        if cursor.rowcount != 1:
                            raise ValueError("profile gained an order during repair; transaction rolled back")
                connection.commit()
                print("committed_updates={} committed_inserts={}".format(len(changes), len(missing_changes)))
            else:
                connection.rollback()
                print("dry_run_only; no database changes")


if __name__ == "__main__":
    main()
